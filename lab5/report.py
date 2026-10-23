"""Отчёт практической работы 5 в оформлении, близком к работе 4."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "results" / "experiment_results.json"
FIG = ROOT / "results" / "figures"
REPO = "https://github.com/shibapc/int-analysis-of-data"


def _font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def _paragraph(doc: Document, text: str, *, size=14, bold=False, align="both", space_before=0, space_after=8, center=False, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center or align == "center" else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    _font(run, size, bold)
    return p


def _heading(doc: Document, text: str) -> None:
    _paragraph(doc, text, size=14, bold=True, space_before=12, space_after=6, keep_with_next=True)


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("– " + text)
    _font(run, 14)


def _hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "28")
    rpr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "28")
    rpr.append(size_cs)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    rpr.append(rfonts)
    new_run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    new_run.append(node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _shade(cell, fill: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell(cell, text: str, *, bold=False, white=False, size=10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    _font(run, size, bold)
    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)
    _shade(cell, "1F4E79" if white else "FFFFFF")


def _table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    table.autofit = True
    for j, name in enumerate(header):
        _set_cell(table.rows[0].cells[j], name, bold=True, white=True)
    for i, row in enumerate(rows):
        fill = "F4F6F8" if i % 2 else "FFFFFF"
        for j, value in enumerate(row):
            _set_cell(table.rows[i + 1].cells[j], value)
            if i % 2:
                _shade(table.rows[i + 1].cells[j], fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def _caption(doc: Document, text: str) -> None:
    _paragraph(doc, text, size=12, align="center", space_before=2, space_after=12, center=True)


def _picture(doc: Document, path: Path, width: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Cm(width))


def _fmt_profile(row: dict) -> list[str]:
    return [
        str(row["cluster"] + 1),
        str(row["n"]),
        f"{row['Age']:.1f}",
        f"{row['Education']:.1f}",
        f"{row['YearsEmployed']:.1f}",
        f"{row['Income']:.1f}",
        f"{row['DebtIncomeRatio']:.1f}",
    ]


def build() -> Path:
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    ks = [str(k) for k in range(3, 10)]
    clean = data["levels"]["0.00"]
    chosen = int(data["chosen_k"])
    km = data["clean_detail"]["kmeans"]
    md = data["clean_detail"]["kmedoids"]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    _paragraph(
        doc,
        "Практическая работа 5. Кластеризация клиентов банка",
        size=16,
        bold=True,
        center=True,
        space_before=12,
        space_after=12,
    )

    _heading(doc, "1. Формулировка задания")
    _paragraph(
        doc,
        "Разработайте программу, которая выполняет кластеризацию заданного набора данных "
        "с помощью алгоритмов k-Means и k-Medoids. Параметрами программы являются набор "
        "данных и число кластеров. Программа должна выдавать координаты точек и назначенные "
        "им кластеры, а также значение ошибки кластеризации.",
    )
    _paragraph(
        doc,
        "Проведите эксперименты на наборе данных customers (сведения о клиентах банка). "
        "Выполните визуализацию: точечный график, на котором цвет точки отражает "
        "принадлежность кластеру, и зависимость ошибки кластеризации от параметра k.",
    )
    _paragraph(
        doc,
        "Доработайте программу, добавив в список её параметров долю зашумленных объектов. "
        "Программа должна случайным образом изменить заданную долю объектов: добавить или "
        "вычесть случайное число у одной или нескольких координат. Проведите эксперименты, "
        "варьируя долю зашумленных объектов (1%, 3%, 5%, 10%) и число кластеров k из "
        "интервала от 3 до 9. Постройте те же диаграммы.",
    )
    _paragraph(
        doc,
        "Подготовьте отчет о выполнении задания и загрузите отчет в формате PDF в систему. "
        "Отчет должен представлять собой связный и структурированный документ со следующими разделами:",
    )
    for item in (
        "формулировка задания;",
        "гиперссылка на каталог репозитория с исходными текстами, наборами данных и др. сопутствующими материалами;",
        "рисунки с результатами визуализации;",
        "пояснения, раскрывающие смысл полученных результатов.",
    ):
        _bullet(doc, item)

    _heading(doc, "2. Гиперссылка на репозиторий")
    p = _paragraph(
        doc,
        "Репозиторий с исходными текстами, набором данных и сопутствующими материалами "
        "(каталог lab5/): ",
        space_after=8,
    )
    # последний run уже содержит текст; добавим ссылку в тот же абзац
    _hyperlink(p, REPO, REPO)

    _heading(doc, "3. Набор данных и метод")
    _paragraph(
        doc,
        "Набор customers содержит 850 клиентов банка. Для расстояния взяты семь признаков: "
        "возраст, код образования (1–5), стаж, годовой доход в тысячах долларов, долг по карте, "
        "прочие долги и отношение долга к доходу. Порядковый номер строки и идентификатор клиента "
        "в расстояние не входят. Признак Defaulted пропущен у 150 клиентов, поэтому столбец "
        "исключён: заполнение пропусков само создало бы искусственную группу.",
    )
    _paragraph(
        doc,
        "Перед кластеризацией каждый признак приводится к нулевому среднему и единичной "
        "дисперсии, иначе доход (размах до 446) подавил бы возраст и стаж. Ошибка k-Means — "
        "сумма квадратов евклидовых расстояний до центроида (SSE). Ошибка k-Medoids — сумма "
        "расстояний до медоида, то есть до реального клиента, ближайшего к остальным в своём "
        "кластере. Оба алгоритма запускаются 10 раз со случайной инициализацией, сохраняется "
        "лучшее разбиение. Точечные графики построены в плоскости двух первых главных компонент "
        f"({data['pca_ratio'][0] * 100:.0f}% и {data['pca_ratio'][1] * 100:.0f}% дисперсии): "
        "в семимерном пространстве кластеры не нарисовать напрямую.",
    )
    _paragraph(
        doc,
        "Шум вносится уже в стандартизованные координаты. У заданной доли клиентов у одной, "
        "двух или трёх случайных координат прибавляется число из равномерного распределения "
        "от −1,5 до 1,5 стандартного отклонения этого признака. Одна и та же зашумленная "
        "матрица подаётся обоим алгоритмам. Запуск: python main.py -k 6; с шумом 5% — "
        "python main.py -k 6 --noise 0.05. Серия экспериментов — python experiments.py.",
    )

    _heading(doc, "4. Сводка экспериментов")
    _paragraph(
        doc,
        "Таблица 1. Ошибка на исходном наборе. Вертикальная линия на рисунках 1 и 2 "
        f"отмечает k = {chosen}: излом кривой SSE, то есть точку, дальше всего отстоящую "
        "от хорды между k = 3 и k = 9. Точечные графики построены для этого k.",
        keep_with_next=True,
    )
    _table(
        doc,
        ["k", "k-Means, SSE", "k-Medoids, сумма расстояний"],
        [
            [
                k,
                f"{clean['kmeans'][k]:.0f}",
                f"{clean['kmedoids'][k]:.0f}",
            ]
            for k in ks
        ],
    )
    _paragraph(
        doc,
        f"Таблица 2. Состав кластеров k-Means при k = {chosen}. Кластеры упорядочены по среднему доходу. "
        "Образование — средний код, доход — тысячи долларов в год, «долг/доход» — проценты.",
    )
    _table(
        doc,
        ["Кластер", "N", "Возраст", "Образование", "Стаж", "Доход", "Долг/доход"],
        [_fmt_profile(row) for row in km["profiles"]],
    )
    _paragraph(
        doc,
        "k-Medoids при том же k выделяет группы того же смысла. Размеры кластеров по возрастанию "
        "дохода: "
        + ", ".join(str(n) for n in md["sizes"])
        + f". Суммарная ошибка медоидов равна {md['error']:.0f}. Медоид — конкретный клиент, "
        "а не среднее координат. У группы из 10 самых обеспеченных клиентов средний доход "
        f"{md['profiles'][-1]['Income']:.0f} тыс. $, а медоид — клиент с доходом "
        f"{md['centers_original'][-1][3]:.0f} тыс. $: среднее утягивают несколько "
        "максимальных доходов, медоид остаётся внутри группы.",
    )

    doc.add_page_break()
    _heading(doc, "5. Визуализация результатов")
    _picture(doc, FIG / "error_kmeans.png", 15.6)
    _caption(doc, "Рисунок 1 – Зависимость SSE алгоритма k-Means от числа кластеров")
    _picture(doc, FIG / "error_kmedoids.png", 15.6)
    _caption(doc, "Рисунок 2 – Зависимость ошибки k-Medoids от числа кластеров")
    _picture(doc, FIG / "scatter_clean.png", 16.2)
    _caption(
        doc,
        f"Рисунок 3 – Кластеры клиентов при k = {chosen}. Цвет точки — номер кластера, "
        "крестик — центр",
    )
    _picture(doc, FIG / "error_noise_kmeans.png", 15.6)
    _caption(doc, "Рисунок 4 – SSE k-Means при доле зашумленных объектов 0–10%")
    _picture(doc, FIG / "error_noise_kmedoids.png", 15.6)
    _caption(doc, "Рисунок 5 – Ошибка k-Medoids при доле зашумленных объектов 0–10%")
    _picture(doc, FIG / "scatter_noise_kmeans.png", 16.2)
    _caption(doc, f"Рисунок 6 – k-Means, k = {chosen}, при доле шума 1%, 3%, 5% и 10%")
    _picture(doc, FIG / "scatter_noise_kmedoids.png", 16.2)
    _caption(doc, f"Рисунок 7 – k-Medoids, k = {chosen}, при доле шума 1%, 3%, 5% и 10%")

    doc.add_page_break()
    _heading(doc, "6. Смысл полученных результатов")
    sse3 = clean["kmeans"]["3"]
    sse6 = clean["kmeans"]["6"]
    sse9 = clean["kmeans"]["9"]
    md3 = clean["kmedoids"]["3"]
    md8 = clean["kmedoids"]["8"]
    md9 = clean["kmedoids"]["9"]
    _paragraph(
        doc,
        f"Ошибка обоих алгоритмов падает с ростом k: у k-Means SSE снижается с {sse3:.0f} "
        f"при трёх кластерах до {sse9:.0f} при девяти, у k-Medoids — с {md3:.0f} до {md9:.0f}. "
        "Так и должно быть: лишний центр всегда может подогнаться к части точек. Содержателен "
        f"не сам минимум, а замедление. У k-Means излом кривой приходится на k = {chosen}: "
        f"от 5 к 6 кластерам SSE ещё теряет {clean['kmeans']['5'] - sse6:.0f}, "
        f"от 6 к 7 — уже {sse6 - clean['kmeans']['7']:.0f}. У k-Medoids спад почти останавливается "
        f"позже: между k = 8 и k = 9 ошибка меняется с {md8:.0f} до {md9:.0f}. "
        "Дальше дробить выборку из 850 клиентов уже мало что добавляет.",
    )
    _paragraph(
        doc,
        "Шесть кластеров k-Means читаются как типы клиентов, а не как случайные куски. "
        "Самая большая группа (336 человек) — молодые клиенты около 31 года со стажем около "
        "5 лет, доходом около 29 тыс. $ и низкой долговой нагрузкой (6,7%). Рядом группа "
        "того же возраста и дохода (139 человек), но с отношением долга к доходу около 19%: "
        "это заёмщики с той же зарплатой и заметно более тяжёлым долгом. Отдельная группа "
        "из 101 человека выделяется образованием (средний код 3,5 при 1,3–1,5 у молодых "
        "групп) и доходом около 42 тыс. $. Четвёртая группа (180 человек) старше, около "
        "42 лет, со стажем 16 лет и доходом 68 тыс. $ при умеренном долге. Пятая (80 человек) "
        "сочетает высокий доход, около 81 тыс. $, с высокой долговой нагрузкой, около 18%. "
        "Шестая — 14 клиентов со средним доходом 213 тыс. $ и крупными долгами: это хвост "
        "распределения, и центроид k-Means уходит к нему сильнее, чем медоид.",
    )
    _paragraph(
        doc,
        "k-Medoids собирает те же шесть смыслов, размеры групп — "
        + ", ".join(str(n) for n in md["sizes"])
        + ". Хвост высокого дохода здесь короче (10 клиентов против 14): медоид обязан "
        "быть реальным клиентом и не сдвигается в пустоту между выбросами. Поэтому "
        "k-Medoids устойчивее к нескольким очень богатым заёмщикам, а k-Means точнее "
        "минимизирует именно сумму квадратов и сильнее реагирует на этот хвост. "
        "На рисунке 3 крайние левые точки — как раз высокий доход и долг: у первой "
        "главной компоненты большие отрицательные веса у дохода, долгов, возраста и стажа. "
        "Вторая компонента отделяет старших клиентов с умеренным долгом (верх) от высокой "
        "долговой нагрузки (низ). Вместе оси держат около 65% дисперсии, поэтому часть "
        "групп, которые различаются в первую очередь образованием, на плоскости пересекаются. "
        "Перекрытие цветов на рисунке — следствие проекции, а не смешения кластеров в исходных семи признаках.",
    )
    noise6_km = data["levels"]["0.10"]["kmeans"]["6"]
    noise6_md = data["levels"]["0.10"]["kmedoids"]["6"]
    noise1_km = data["levels"]["0.01"]["kmeans"]["6"]
    _paragraph(
        doc,
        f"Шум до 5% почти не меняет кривые. При k = {chosen} SSE k-Means составляет "
        f"{sse6:.0f} без шума, {noise1_km:.0f} при 1% и {data['levels']['0.05']['kmeans']['6']:.0f} "
        f"при 5%. Десять процентов зашумленных объектов поднимают SSE до {noise6_km:.0f} "
        f"(+{(noise6_km - sse6) / sse6 * 100:.1f}%) и ошибку k-Medoids с {md['error']:.0f} "
        f"до {noise6_md:.0f} (+{(noise6_md - md['error']) / md['error'] * 100:.1f}%). "
        "Форма зависимости от k сохраняется. На рисунке 4 линия 10% лежит выше остальных "
        "на всём интервале, излом не переезжает. На рисунке 5 сдвиг меньше, и отдельные "
        "кривые пересекаются: медоид слабее реагирует на сдвиг нескольких координат, "
        "а один случайный набор шума даёт колебания в несколько десятков единиц. "
        "На рисунках 6 и 7 при всех четырёх долях шума видна та же "
        "картина, что и на чистых данных: плотное облако справа, группа высокого дохода "
        "слева, долговая группа ниже центра. Для этой выборки шум в пределах 10% объектов "
        "с амплитудой полутора стандартных отклонений не разрушает найденные типы клиентов. "
        "Менять решение о числе кластеров из-за такого шума нет причин: и при чистых данных, "
        "и при шуме разумный диапазон — от 4 до 6 групп, а дробление дальше восьми почти "
        "не уменьшает ошибку k-Medoids.",
    )

    out = ROOT / "Отчет по практической работе 5.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(path)
