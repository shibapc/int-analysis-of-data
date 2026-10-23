"""Отчёт практической работы 6 в оформлении работ 4 и 5."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "results" / "experiment_results.json"
FIG = ROOT / "results" / "figures"
REPO = "https://github.com/shibapc/int-analysis-of-data"
ORDER = ("single", "complete", "average", "ward")


def _font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def _paragraph(doc, text, *, size=14, bold=False, align="both", space_before=0, space_after=8, center=False, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center or align == "center" else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    _font(run, size, bold)
    return p


def _heading(doc, text: str) -> None:
    _paragraph(doc, text, size=14, bold=True, space_before=12, space_after=6, keep_with_next=True)


def _bullet(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("– " + text)
    _font(run, 14)


def _hyperlink(paragraph, text: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "28")
    rpr.append(size)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rpr.append(rfonts)
    new_run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    new_run.append(node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell(cell, text: str, *, bold=False, white=False, size=10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    _font(run, size, bold)
    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)
    _shade(cell, "1F4E79" if white else "FFFFFF")


def _table(doc, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for j, name in enumerate(header):
        _set_cell(table.rows[0].cells[j], name, bold=True, white=True, size=9)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            _set_cell(table.rows[i + 1].cells[j], value, size=10)
            if i % 2:
                _shade(table.rows[i + 1].cells[j], "F4F6F8")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def _caption(doc, text: str) -> None:
    _paragraph(doc, text, size=12, center=True, space_before=2, space_after=12)


def _picture(doc, path: Path, width: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Cm(width))


def _sizes(values: list[int]) -> str:
    return ", ".join(str(n) for n in sorted(values, reverse=True))


def _num(value: float) -> str:
    return f"{value:.2f}".replace(".", ",")


def build() -> Path:
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    methods = data["methods"]
    ward = methods["ward"]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    _paragraph(
        doc,
        "Практическая работа 6. Иерархическая кластеризация",
        size=16,
        bold=True,
        center=True,
        space_before=12,
        space_after=12,
    )

    _heading(doc, "1. Формулировка задания")
    _paragraph(
        doc,
        "Разработайте программу, которая выполняет иерархическую кластеризацию заданного "
        "набора данных. Параметрами программы являются набор данных и используемая мера "
        "схожести кластеров (Single linkage, Complete linkage, Group average или Ward distance). "
        "Мера схожести объектов произвольная (в этой работе зафиксировано евклидово расстояние). "
        "Программа должна выполнять визуализацию полученных результатов в виде дендрограммы.",
    )
    _paragraph(
        doc,
        "Проведите эксперименты на наборе данных из задания «Разделительная кластеризация» "
        "(customers, клиенты банка). Сравните полученные результаты с результатами k-Means "
        "и k-Medoids. Проведите эксперименты, варьируя меру схожести кластеров, и сравните "
        "эти меры между собой.",
    )
    _paragraph(
        doc,
        "Подготовьте отчет о выполнении задания и загрузите отчет в формате PDF в систему. "
        "Отчет должен представлять собой связный и структурированный документ со следующими разделами:",
    )
    for item in (
        "формулировка задания;",
        "гиперссылка на каталог репозитория с исходными текстами, наборами данных и др. сопутствующими материалами;",
        "рисунки с результатами визуализации;",
        "пояснения, раскрывающие смысл полученных результатов.",
    ):
        _bullet(doc, item)

    _heading(doc, "2. Гиперссылка на репозиторий")
    p = _paragraph(
        doc,
        "Репозиторий с исходными текстами, набором данных и сопутствующими материалами "
        "(каталог lab6/): ",
    )
    _hyperlink(p, REPO, REPO)

    _heading(doc, "3. Набор данных и метод")
    _paragraph(
        doc,
        "Набор и подготовка те же, что в практической работе 5: 850 клиентов, семь признаков "
        "(возраст, код образования, стаж, доход, долг по карте, прочие долги, отношение долга "
        "к доходу). Идентификатор клиента отброшен. Столбец Defaulted исключён из-за 150 пропусков. "
        "Признаки стандартизованы, иначе доход подавил бы остальные координаты. Расстояние между "
        "объектами везде евклидово.",
    )
    _paragraph(
        doc,
        "Дерево строится снизу вверх. Single linkage объединяет кластеры по ближайшей паре точек, "
        "Complete linkage — по самой дальней, Group average — по среднему расстоянию всех пар, "
        "Ward — по приросту внутрикластерной дисперсии. На дендрограмме показаны верхние слияния "
        "(нижние 850 листьев схлопнуты в скобки с числом клиентов, иначе рисунок был бы нечитаем). "
        "Пунктир — высота разреза на 6 кластеров: столько групп было выбрано в работе 5 по излому SSE. "
        "Высоты разных мер нельзя сравнивать между панелями: у Ward на оси не евклидово расстояние, "
        "а прирост дисперсии. Запуск: python main.py -m ward. Все четыре меры — python experiments.py.",
    )

    _heading(doc, "4. Сводка экспериментов")
    _paragraph(
        doc,
        "Таблица 1. Разрез каждого дерева на 6 кластеров. Размеры даны по убыванию. "
        "Согласованность с работой 5 измерена скорректированным индексом Рэнда: 1 — "
        "разбиения совпадают, 0 — совпадение не лучше случайного.",
        keep_with_next=True,
    )
    _table(
        doc,
        ["Мера", "Размеры 6 кластеров", "С k-Means", "С k-Medoids"],
        [
            [
                methods[name]["label"],
                _sizes(methods[name]["sizes"]),
                _num(methods[name]["ari_kmeans"]),
                _num(methods[name]["ari_kmedoids"]),
            ]
            for name in ORDER
        ],
    )
    _paragraph(
        doc,
        "Для сравнения: k-Means в работе 5 при k = 6 дал размеры 336, 180, 139, 101, 80, 14, "
        "k-Medoids — 271, 188, 140, 134, 107, 10 (порядок — по убыванию).",
    )
    _paragraph(
        doc,
        "Таблица 2. Состав кластеров Ward при том же разрезе, кластеры упорядочены по среднему доходу.",
        keep_with_next=True,
    )
    _table(
        doc,
        ["Кластер", "N", "Возраст", "Образование", "Стаж", "Доход", "Долг/доход"],
        [
            [
                str(row["cluster"] + 1),
                str(row["n"]),
                f"{row['Age']:.1f}",
                f"{row['Education']:.1f}",
                f"{row['YearsEmployed']:.1f}",
                f"{row['Income']:.1f}",
                f"{row['DebtIncomeRatio']:.1f}",
            ]
            for row in ward["profiles"]
        ],
    )

    _heading(doc, "5. Визуализация результатов")
    _picture(doc, FIG / "dendrograms.png", 16.2)
    _caption(doc, "Рисунок 1 – Дендрограммы четырёх мер. Пунктир — разрез на 6 кластеров")
    _picture(doc, FIG / "dendrogram_single.png", 16.0)
    _caption(doc, "Рисунок 2 – Single linkage: цепочечное слияние и отделение единичных объектов")
    _picture(doc, FIG / "dendrogram_ward.png", 16.0)
    _caption(doc, "Рисунок 3 – Ward: ветви сопоставимого размера")
    _picture(doc, FIG / "scatter_k6.png", 16.2)
    _caption(
        doc,
        "Рисунок 4 – Те же клиенты в плоскости главных компонент работы 5. "
        "Цвет — кластер после разреза дендрограммы на 6 групп",
    )

    _heading(doc, "6. Смысл полученных результатов")
    _paragraph(
        doc,
        "Четыре меры на одном и том же расстоянии дают разные деревья. Single linkage "
        f"при разрезе на 6 кластеров оставляет {methods['single']['sizes'][0]} клиентов "
        "в одной группе и выносит наружу несколько самых далёких точек: три клиента "
        "с доходом около 128 тыс. $ и тяжёлым долгом и четырёх одиночек, среди которых "
        "клиент с доходом 446 тыс. $. На рисунке 2 это видно как лестница: каждый следующий "
        "шаг цепляется за ближайшую точку и почти не поднимает порог, пока не остаются "
        "выбросы. Group average ведёт себя так же: "
        f"{methods['average']['sizes'][0]} клиентов в главной группе, остальные кластеры — "
        "от 1 до 9 человек. Индекс Рэнда с k-Means у обеих мер около 0,01–0,02: общего "
        "разбиения с работой 5 нет, иерархия здесь только выделяет хвост.",
    )
    _paragraph(
        doc,
        "Complete linkage смотрит на самую дальнюю пару, поэтому цепочка обрывается раньше. "
        "Шесть кластеров — это 555 молодых клиентов с доходом около 30 тыс. $, 260 более "
        "старших со стажем около 15 лет и доходом около 74 тыс. $, группа из 23 человек "
        "с долговой нагрузкой около 26% и три мелкие группы с доходом от 114 до 305 тыс. $. С k-Means "
        f"согласованность уже {_num(methods['complete']['ari_kmeans'])}: крупные куски "
        "«молодых» и «старших» узнаются, но дробление по долгу и образованию, которое "
        "сделали k-Means и k-Medoids, complete linkage при шести кластерах ещё не выделяет.",
    )
    rich = ward["profiles"][-1]
    young_low = ward["profiles"][1]
    young_debt = ward["profiles"][0]
    educated = ward["profiles"][2]
    older = ward["profiles"][4]
    debt_rich = ward["profiles"][3]
    _paragraph(
        doc,
        f"Ward ближе всего к разделительной кластеризации: индекс Рэнда с k-Means "
        f"{_num(ward['ari_kmeans'])}, с k-Medoids {_num(ward['ari_kmedoids'])}. "
        "Состав групп повторяет работу 5. Самая большая группа — "
        f"{young_low['n']} человек около {young_low['Age']:.0f} лет, доход около "
        f"{young_low['Income']:.0f} тыс. $ и долг/доход {young_low['DebtIncomeRatio']:.1f}% "
        f"(в k-Means такая группа была из 336 человек). Рядом {young_debt['n']} молодых "
        f"с той же зарплатой и долговой нагрузкой {young_debt['DebtIncomeRatio']:.1f}% "
        f"(в k-Means — 139 человек и 19%). Отдельно {educated['n']} клиентов с кодом "
        f"образования {educated['Education']:.1f} и доходом около {educated['Income']:.0f} тыс. $ "
        f"(в k-Means — 101 человек и код 3,5). Старшие со стажем {older['YearsEmployed']:.0f} лет "
        f"и умеренным долгом — {older['n']} человек; обеспеченные с нагрузкой "
        f"{debt_rich['DebtIncomeRatio']:.0f}% — {debt_rich['n']} человек. Хвост высокого дохода — "
        f"{rich['n']} клиентов со средним доходом {rich['Income']:.0f} тыс. $ "
        "(в k-Means их было 14, в k-Medoids — 10). На рисунке 4 только панель Ward "
        "раскрашивает плотное облако так же, как точечный график работы 5; single linkage "
        "и group average оставляют облако одним цветом и красят отдельные точки слева.",
    )
    _paragraph(
        doc,
        "Расхождение объясняется тем, что минимизируют разные величины. k-Means и Ward "
        "уменьшают разброс внутри группы, поэтому ищут компактные типы клиентов. "
        "k-Medoids делает то же на реальных объектах и потому чуть дальше от Ward "
        f"(индекс {_num(ward['ari_kmedoids'])} против {_num(ward['ari_kmeans'])} с k-Means): "
        "медоид не уезжает в хвост дохода. Single linkage и group average минимизируют "
        "локальную близость, и на выборке с несколькими очень богатыми заёмщиками "
        "дерево тратит разрезы на эти точки, а не на типы. Для customers из работы 5 "
        "содержательное разбиение даёт Ward; single linkage полезен, если задача — "
        "найти единичные необычные анкеты, а не сегменты.",
    )

    out = ROOT / "Отчет по практической работе 6.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build())
